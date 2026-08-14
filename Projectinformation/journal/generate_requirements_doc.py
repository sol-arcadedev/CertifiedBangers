"""
Generates Requirements_Engineering.docx — the consolidated project requirements
specification for CertifiedBanger, synthesized from README.md and the
Development Journal (entries.json).

Usage: python generate_requirements_doc.py
"""
import os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = os.path.dirname(os.path.abspath(__file__))
OUTPUT_PATH = os.path.join(os.path.dirname(HERE), "Requirements_Engineering.docx")


# ---------- low-level helpers (shared technique with generate_journal.py) ----------

def add_field(paragraph, instr_text, display_text):
    run = paragraph.add_run()
    r = run._r
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instr_text
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    text_el = OxmlElement("w:t")
    text_el.text = display_text
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    r.append(fld_begin); r.append(instr); r.append(fld_sep); r.append(text_el); r.append(fld_end)


def enable_auto_update_fields(document):
    settings = document.settings.element
    update_fields = OxmlElement("w:updateFields")
    update_fields.set(qn("w:val"), "true")
    settings.append(update_fields)


def build_styles(document):
    styles = document.styles

    title_style = styles.add_style("DocTitle", WD_STYLE_TYPE.PARAGRAPH)
    title_style.base_style = styles["Title"]
    title_style.font.size = Pt(28)
    title_style.font.color.rgb = RGBColor(0x1A, 0x1A, 0x2E)

    meta_style = styles.add_style("Meta", WD_STYLE_TYPE.PARAGRAPH)
    meta_style.font.size = Pt(10)
    meta_style.font.color.rgb = RGBColor(0x60, 0x60, 0x60)
    meta_style.font.italic = True

    code_style = styles.add_style("CodeBlock", WD_STYLE_TYPE.PARAGRAPH)
    code_style.font.name = "Consolas"
    code_style.font.size = Pt(9)
    code_style.font.color.rgb = RGBColor(0x20, 0x20, 0x20)


# ---------- content-block renderer ----------

def render_blocks(document, blocks):
    for block in blocks:
        kind = block[0]
        if kind == "h":
            _, level, text = block
            document.add_heading(text, level=level)
        elif kind == "p":
            _, text = block
            document.add_paragraph(text)
        elif kind == "bullets":
            _, items = block
            for item in items:
                document.add_paragraph(item, style="List Bullet")
        elif kind == "numbered":
            _, items = block
            for item in items:
                document.add_paragraph(item, style="List Number")
        elif kind == "table":
            _, headers, rows = block
            table = document.add_table(rows=1, cols=len(headers))
            table.style = "Light Grid Accent 1"
            hdr = table.rows[0].cells
            for i, h in enumerate(headers):
                hdr[i].text = h
                for para in hdr[i].paragraphs:
                    for run in para.runs:
                        run.bold = True
            for row_data in rows:
                row = table.add_row().cells
                for i, val in enumerate(row_data):
                    row[i].text = val
            document.add_paragraph()
        elif kind == "code":
            _, text = block
            for line in text.split("\n"):
                p = document.add_paragraph(style="CodeBlock")
                p.paragraph_format.space_after = Pt(0)
                p.add_run(line if line else " ")
        elif kind == "pagebreak":
            document.add_page_break()
        elif kind == "spacer":
            document.add_paragraph()


# ---------- document content ----------

def add_title_page(document):
    p = document.add_paragraph(style="DocTitle")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("CertifiedBanger")

    sub = document.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = sub.add_run("Project Requirements Engineering Document")
    run.font.size = Pt(18)
    run.font.color.rgb = RGBColor(0x40, 0x40, 0x40)

    from datetime import datetime
    meta = document.add_paragraph(style="Meta")
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.add_run(f"Generated {datetime.now().strftime('%Y-%m-%d')} · v1.0")

    document.add_paragraph()
    purpose = document.add_paragraph()
    purpose.add_run(
        "This document consolidates the CertifiedBanger project brief (README.md) and every "
        "decision recorded in the Development Journal (Development_Journal.docx) into a single, "
        "current requirements specification. It describes the project, its functional and "
        "non-functional requirements, security considerations, the chosen technology stack, "
        "a possible software architecture, the data model, a phased work-package build plan, "
        "and a user-story backlog to work from task by task. Where a requirement traces back to "
        "a specific journal decision, the entry number is referenced (e.g. 'Journal Entry 12') "
        "so the full rationale can be looked up if needed — this document states the *what*, "
        "the journal holds the *why*."
    )
    document.add_page_break()


def add_toc(document):
    document.add_heading("Table of Contents", level=1)
    toc_p = document.add_paragraph()
    add_field(
        toc_p,
        'TOC \\o "1-2" \\h \\z \\u',
        "Table of Contents will appear here — in Word, right-click and choose "
        "'Update Field', or press Ctrl+A then F9, to generate it.",
    )
    document.add_page_break()


SECTION_1 = [
    ("h", 1, "1. Introduction & Project Overview"),
    ("h", 2, "1.1 Vision"),
    ("p", "CertifiedBanger is a community-curated review platform where readers rate and review "
          "manga/manhwa across multiple structured quality dimensions, and the community can award "
          "the best entries a distinct, meaningful quality seal. Version 1 focuses exclusively on "
          "manga/manhwa; the architecture is deliberately designed so a future expansion to other "
          "consumable media (games, movies, novels) is a schema extension, not a rewrite."),
    ("p", "Manga/manhwa readers currently rely on scattered, low-signal sources — general databases "
          "with a single aggregate score, forum threads, social media — to find their next read. "
          "CertifiedBanger's core value proposition: instead of one opaque score, users see *why* "
          "something is good, broken into consistent categories, backed by a review, and — when a "
          "title is exceptional — marked with a community-recognized seal of quality."),
    ("h", 2, "1.2 Goals (v1)"),
    ("bullets", [
        "Let users write structured reviews with per-category ratings.",
        "Let the community award a distinct 'seal' (Certified Banger, Hidden Gem) to reviews they believe deserve special recognition, backed by a verification workflow that protects the seal's credibility.",
        "Let users rate/comment on other users' reviews, surfacing the best reviews, not just the best titles.",
        "Provide browsing/discovery: search, filter, sort by category score, seal count, etc.",
        "Let users track their own reading progress via a personal library.",
        "Build the data model so a future generalization to other media types is a schema extension, not a rewrite.",
    ]),
    ("h", 2, "1.3 Non-Goals (v1)"),
    ("bullets", [
        "No token/crypto integration in v1 (planned for a later version once the reward is tied to real usage — see Section 2.9).",
        "No manga/manhwa reading functionality — this is not a scanlation/hosting site. The platform reviews titles; it does not distribute them.",
        "No support for media types other than manga/manhwa in v1.",
        "No native mobile app in v1 (responsive web only).",
        "No monetization/ads in v1.",
        "No fully-automatic recommendation engine in v1 — the underlying library data ships in v1, but the recommendation algorithm itself is future scope (Journal Entry 20).",
    ]),
    ("h", 2, "1.4 Target Users"),
    ("bullets", [
        "Primary persona: an engaged manga/manhwa reader who has read widely, has strong opinions, and is frustrated by generic single-score databases.",
        "Secondary persona: a casual reader looking for their next read, who wants trustworthy, structured, skimmable recommendations rather than a single number.",
    ]),
]

SECTION_2 = [
    ("h", 1, "2. Functional Requirements"),
    ("h", 2, "2.1 Titles (Catalog)"),
    ("p", "A Title represents a single manga/manhwa/manhua work (the series as a whole, not a chapter "
          "or volume)."),
    ("bullets", [
        "Fields: name, alternate/original names, type (manga/manhwa/manhua — extensible enum), author(s)/illustrator(s), status (ongoing/completed/hiatus/dropped), genre tags, cover image, publication year, source links (official reading platforms only — never hosted content), short synopsis.",
        "Titles are de-duplicated: search-before-create UX for users/admins, plus admin merge tooling for canonical entries.",
        "Titles aggregate and display: average score per category, seal counts (Certified Banger, Hidden Gem), review count, and the review list.",
        "For performance at scale, these aggregates are stored as precomputed columns updated incrementally, not recalculated live on every page view (Journal Entry 39).",
        "Admin-seeded at launch — no scraping for v1 (matches the original MVP cut).",
    ]),
    ("h", 2, "2.2 Reviews"),
    ("bullets", [
        "One review per user per title; editing replaces the existing review rather than creating a duplicate.",
        "A numeric rating per category, using the platform's data-driven Category set (Art Style, Character, Plot, Pacing, Uniqueness for v1) — Journal Entry 2.",
        "Overall score is computed automatically as the simple average of the five category scores — not a separately entered field (Journal Entry 4).",
        "Free-text written review body (recommended minimum ~50 characters to discourage low-effort spam), with an optional spoiler flag/section.",
        "Reviews track created_at / updated_at.",
        "Reviews can be upvoted/downvoted by other users, and can receive flat (non-threaded) comments.",
        "A baseline gate applies to submitting any review, active from launch: email verification plus a minimum account age of 3 days. No specific future restrictions are pre-committed — the Main Admin adjusts the gate's (admin-configurable) parameters over time based on observed abuse signals (Journal Entries 16, 40).",
        "Admin-authored reviews are visually labeled '<username> (admin)' everywhere the username appears on the platform — not just on the review itself (Journal Entry 27).",
        "Reviews written by an admin other than the Main Admin require the Main Admin's approval before they publish. The Main Admin's own reviews publish immediately. This requirement is lifted for a given admin at the Main Admin's manual, case-by-case discretion — no automatic rule (Journal Entries 28, 41). A pending admin review is fully hidden — not visible, votable, commentable, or seal-eligible — until approved (Journal Entry 42).",
    ]),
    ("h", 2, "2.3 Seals — Certified Banger & Hidden Gem"),
    ("p", "Seals are data-driven (a SealType table), not hardcoded, so new seal types can be added "
          "later without a code change. Two seal types ship in v1: Certified Banger (the flagship "
          "seal for exceptional quality) and Hidden Gem (strong quality, low popularity — a "
          "discovery tool for under-the-radar titles). A seal attaches to a specific review, not "
          "the title directly; a title's displayed seal count is the number of its reviews that "
          "have earned that seal (Journal Entry 7)."),
    ("h", 3, "2.3.1 Verification Workflow"),
    ("p", "Seal-granting is not a self-service user action. It follows a two-phase verification "
          "workflow, which applies generically to every seal type (Journal Entries 6, 14):"),
    ("bullets", [
        "Phase 1 (launch): the Main Admin (or a trusted admin, per Section 2.4) manually reviews a submitted review and grants the seal after verification.",
        "Phase 2 (post-growth): once a review's community vote net score (upvotes − downvotes) reaches a configurable threshold — set to +20 for launch — it becomes an automatic seal candidate (Journal Entries 8, 29).",
        "To be considered at all, a seal nomination requires the user to have written a full review, plus a separate written justification field distinct from the review body, explaining why the title deserves the seal (Journal Entry 5).",
    ]),
    ("h", 3, "2.3.2 Probation & Permanence"),
    ("bullets", [
        "Probation applies only to Phase 2 (automatic, vote-threshold-granted) seals. Phase 1 admin-granted seals are permanent immediately (Journal Entry 11).",
        "A Phase 2 seal starts 'provisional.' Its net vote score must remain non-negative every single day; any dip resets a streak counter to zero. There is no outright revocation for failing probation — the review simply stays provisional indefinitely until it achieves 30 consecutive positive days, at which point it converts to permanent (Journal Entry 12).",
        "Once permanent, a seal is locked against automatic, vote-based revocation forever. A human admin retains a manual override to demote a permanent seal back to provisional or remove it entirely (Journal Entry 13).",
    ]),
    ("h", 3, "2.3.3 Hidden Gem Criteria"),
    ("bullets", [
        "Quality gate: same positive-vote bar as Certified Banger (net score ≥ the configured threshold).",
        "Popularity gate: measured as total vote count on the title's highest-voted review — set to 100 votes for launch (Journal Entry 29). Below this, a qualifying review earns Hidden Gem; at or above it, the title also earns Certified Banger.",
        "A title can hold both Hidden Gem and Certified Banger simultaneously — earning Certified Banger does not remove Hidden Gem. Hidden Gem remains permanently attached once earned, even after the title later becomes popular, so it stays an accurate historical/discovery marker (Journal Entry 15).",
    ]),
    ("h", 2, "2.4 Users, Profiles & Admin Roles"),
    ("bullets", [
        "Registration/auth: email + password minimum; social login is a nice-to-have.",
        "Anonymous visitors can browse titles, read reviews, and view the seal showcase without an account. An account is required only to write, rate, vote, comment, or take any seal-related action (Journal Entry 3).",
        "Public profile: username, avatar, bio, list of reviews, seals awarded, join date, a basic reputation indicator, and the user's library (Section 2.5).",
        "Two-tier admin structure: a single Main Admin (the platform founder/owner) and additional Admin accounts recruited later. Only the Main Admin can grant/lift the review-approval requirement for a given admin (Journal Entry 28).",
        "Follow/unfollow other users — nice-to-have, not core-critical for v1.",
    ]),
    ("h", 2, "2.5 User Library (Reading Status Tracking)"),
    ("p", "A v1 feature (introduced this session, not in the original brief): registered users can "
          "add titles to a personal library on their profile."),
    ("bullets", [
        "Four statuses: Finished, Currently Reading, Plan to Read, Dropped (Journal Entry 18).",
        "Library entries are fully independent of reviews — a user can add any title at any status with or without ever reviewing it (Journal Entry 19).",
        "This feature ships in v1; it is the data foundation for the future recommendation engine, which is itself deferred (Journal Entry 20).",
    ]),
    ("h", 2, "2.6 Voting & Comments"),
    ("bullets", [
        "Upvote/downvote (or 'helpful' mark) on reviews.",
        "Flat (non-threaded) comments on reviews, for discussion or disagreement.",
        "Vote tallies on a review directly drive Phase 2 seal candidacy and probation tracking (Section 2.3).",
    ]),
    ("h", 2, "2.7 Discovery & Browsing"),
    ("bullets", [
        "Search titles by name, tag, genre, or author (PostgreSQL native full-text search — Journal Entry 38).",
        "Filter by genre, status, category-score thresholds, seal presence, seal count.",
        "Sort by highest overall score, most seals, most recent reviews, most discussed.",
        "A dedicated Certified Banger / Hidden Gem showcase page — the platform's core differentiator, so it should be prominent.",
    ]),
    ("h", 2, "2.8 Moderation & Reporting"),
    ("bullets", [
        "Report flow for reviews, comments, and titles.",
        "Admin panel: view reports, remove content, ban/suspend users, edit title metadata, manage the Category and SealType tables without a code deploy.",
        "Rate limiting / spam prevention on review and comment creation.",
    ]),
    ("h", 2, "2.9 Future-Scope Functional Items (Not v1)"),
    ("p", "Documented now for architectural awareness, per the brief's own approach to future scope "
          "(Section 8) and this session's additions:"),
    ("bullets", [
        "Personalized recommendation engine — content-based, using a user's category-rating patterns and genre affinity from 'Finished' library entries only. Algorithm design intentionally deferred until the feature enters active development (Journal Entries 21, 24, 25).",
        "First-review token reward — if the first-ever review of a given title (platform-wide, any user) is well received by vote, its author receives the existing token-reward system's reward. Exact 'well received' threshold not yet defined (Journal Entries 22, 26).",
        "Token/crypto rewards for active contributors (pump.fun-launched coin), deferred until the core platform has real usage; flagged for future legal/compliance review.",
        "Multi-media expansion, public API, reviewer reputation/leveling system, native mobile app — per the original brief's Section 8.",
    ]),
]

SECTION_3 = [
    ("h", 1, "3. Non-Functional Requirements"),
    ("h", 2, "3.1 Performance & Scalability"),
    ("bullets", [
        "Title pages and browse/search should feel instant; category-based sorting/filtering must stay efficient as the catalog grows.",
        "Aggregate values (average category scores, seal counts) are precomputed and updated incrementally, not recalculated live on read (Journal Entry 39).",
        "Composite indexes are built to match the platform's actual filter/sort access patterns (genre + score sort, seal-presence + seal-count sort), not just foreign keys.",
        "Read replicas are the identified scaling lever if recommendation/analytics-style queries begin competing with normal traffic — an infrastructure upgrade, not a re-platform (Journal Entry 39).",
    ]),
    ("h", 2, "3.2 Data Integrity & Extensibility"),
    ("bullets", [
        "Rating categories and seal types are centrally defined in the database (Category, SealType tables), not duplicated across the codebase — adding a 6th category or a new seal type is a data change, not a deploy (Journal Entries 2, 14).",
        "The Title / Review / Category schema is designed so a future media type can reuse the same Review/Seal/Category infrastructure rather than requiring parallel tables.",
    ]),
    ("h", 2, "3.3 Abuse Resistance"),
    ("bullets", [
        "The seal's credibility is the core product asset — review-bombing, sockpuppet seal-awarding, and spam must be resisted from day one, not bolted on later.",
        "The phased admin-then-vote verification workflow, provisional probation period, and permanent-seal lock (Section 2.3) exist specifically to make the seal hard to game.",
        "A baseline account-age gate on review submission (Journal Entry 16), rate limiting on review/comment creation, and report tooling provide baseline spam/sockpuppet resistance.",
    ]),
    ("h", 2, "3.4 Content Legality"),
    ("bullets", [
        "No hosting of copyrighted manga/manhwa content or scans — only metadata, cover art, and links to official sources.",
        "Cover art usage should be verified against a fair-use/promotional-norms licensing approach.",
    ]),
    ("h", 2, "3.5 Accessibility"),
    ("bullets", [
        "Basic a11y from launch: alt text on covers, keyboard navigation, sufficient color contrast — not exhaustive for v1, but not ignored either.",
    ]),
    ("h", 2, "3.6 Responsiveness"),
    ("bullets", [
        "Must work well on mobile browsers — this audience browses heavily on phones. Server-side rendering (Section 5) also supports fast perceived load times on mobile networks.",
    ]),
]

SECTION_4 = [
    ("h", 1, "4. Security Considerations"),
    ("h", 2, "4.1 Authentication & Session Security"),
    ("bullets", [
        "Authentication is delegated to Supabase Auth (Journal Entry 34) rather than a hand-rolled implementation — password hashing, session/token management, and password-reset flows are handled by a maintained, audited service rather than custom code.",
        "Sessions should use secure, HTTP-only cookies; no auth tokens stored in client-accessible storage where avoidable.",
    ]),
    ("h", 2, "4.2 Authorization / Role-Based Access Control"),
    ("bullets", [
        "Three roles: regular User, Admin, Main Admin (Journal Entry 28). Every admin-only action (title editing, review approval, seal verification, moderation, report handling, Category/SealType configuration) must check role server-side — never trust a client-side role flag.",
        "The admin-authored review-approval gate (Section 2.2) and the manual seal-permanence override (Section 2.3.2) are both Main-Admin-only actions and must be enforced at the API/server-action layer, not just hidden in the UI.",
        "Anonymous vs. authenticated access boundaries (Journal Entry 3) — read-only browsing endpoints stay public; every write action (review, vote, comment, seal nomination, library entry) requires an authenticated session, checked server-side.",
    ]),
    ("h", 2, "4.3 Input Validation & Injection Prevention"),
    ("bullets", [
        "All database access goes through Prisma's parameterized queries (Journal Entry 33) — no raw string-concatenated SQL, which eliminates the standard SQL-injection vector by construction.",
        "Free-text fields (review body, justification text, comments, bio) must be sanitized/escaped on render to prevent stored XSS, since this content is displayed to other users.",
        "Server-side validation of all inputs (category score ranges, required justification length, spoiler flag, file types/sizes for cover uploads) — client-side validation is a UX convenience only, never the security boundary.",
    ]),
    ("h", 2, "4.4 Abuse & Spam Mitigation"),
    ("bullets", [
        "Rate limiting on review, comment, vote, and report creation (Section 3.3).",
        "Baseline account-age gate before a user can submit their first review (Journal Entry 16) — mitigates disposable-account spam at the entry point to the whole review pipeline.",
        "Report flow with an admin review queue, so user-flagged abuse gets a human response path.",
    ]),
    ("h", 2, "4.5 Seal-Gaming / Vote-Manipulation Resistance"),
    ("bullets", [
        "This is the platform's highest-value security concern, since seal credibility is described repeatedly in the brief as the core product asset.",
        "Defense in depth: (1) a human admin gate at launch when vote volume is too low to trust algorithmically (Journal Entry 6); (2) a net-score threshold, not a raw vote count, so mass single-direction voting is somewhat blunted; (3) a 30-consecutive-day provisional streak that resets on any negative dip, resisting short-lived vote brigades (Journal Entry 12); (4) a manual admin override that can strip even a 'permanent' seal if fraud is discovered later (Journal Entry 13).",
        "One vote per user per review should be enforced at the database level (unique constraint on user_id + target_id + target_type), not just in application logic, to prevent trivial repeat-voting.",
        "Consider IP/device heuristics or account-age weighting on votes as a future hardening step once real abuse patterns are observed — not required for v1, but worth monitoring for.",
    ]),
    ("h", 2, "4.6 Infrastructure & Secrets Management"),
    ("bullets", [
        "Supabase connection strings, Prisma database URLs, and any API keys are stored as environment variables in Vercel's encrypted project settings — never committed to the repository.",
        "HTTPS is enforced by default on Vercel; no plaintext HTTP endpoints.",
        "Cover-image uploads go through Supabase Storage with server-side validation of file type/size before acceptance, to avoid arbitrary file upload risks.",
    ]),
    ("h", 2, "4.7 Content & Legal Risk"),
    ("bullets", [
        "No hosting of copyrighted scans/chapters (Section 3.4) removes the platform's largest legal exposure by design.",
        "Report tooling and admin takedown capability provide a response path for any content dispute (e.g. a disputed cover image).",
    ]),
]

TECH_TABLE_ROWS = [
    ("Framework", "Next.js (TypeScript, App Router)", "Journal Entry 31"),
    ("Database", "PostgreSQL", "Journal Entry 32"),
    ("Database/Auth/Storage host", "Supabase", "Journal Entries 32, 34, 35"),
    ("ORM", "Prisma", "Journal Entry 33"),
    ("Authentication", "Supabase Auth", "Journal Entry 34"),
    ("Cover image storage", "Supabase Storage (Cloudflare R2 migration path)", "Journal Entry 35"),
    ("Background/scheduled jobs", "Vercel Cron", "Journal Entry 36"),
    ("Application hosting", "Vercel (Hobby tier at launch)", "Journal Entry 37"),
    ("Search", "PostgreSQL native full-text search (tsvector/pg_trgm)", "Journal Entry 38"),
    ("Recommendation-readiness", "pgvector extension (path only, not yet used) + precomputed aggregates + read replicas", "Journal Entry 39"),
]

SECTION_5 = [
    ("h", 1, "5. Technology Stack"),
    ("p", "Selected for a solo developer working in the JavaScript/TypeScript ecosystem, prioritizing "
          "minimal operational overhead and a free/cheap launch with a clear upgrade path as traction "
          "grows (Journal Entry 30). Full evaluation and alternatives considered for each layer are "
          "recorded in the Development Journal."),
    ("table", ["Layer", "Choice", "Reference"], TECH_TABLE_ROWS),
]

ARCHITECTURE_DIAGRAM = """
+-----------------------------------+
|              Browser               |
|      (React UI via Next.js)        |
+-------------------+---------------+
                    | HTTPS
                    v
+-----------------------------------+
|        Next.js App (Vercel)        |
|  - SSR pages (SEO: titles, reviews,|
|    Certified Banger/Hidden Gem     |
|    showcase)                       |
|  - API routes / Server Actions     |
|    (reviews, votes, seals, library,|
|    admin actions)                  |
|  - Vercel Cron                     |
|    (daily probation-streak job,    |
|     Journal Entry 36)              |
+-------------------+---------------+
                    | Prisma (Postgres wire protocol)
                    v
+-----------------------------------+
|              Supabase              |
|  - PostgreSQL (all entities,       |
|    Section 7; precomputed          |
|    aggregates, Journal Entry 39)   |
|  - Auth (sessions, roles)          |
|  - Storage (cover images)          |
+-----------------------------------+
""".strip("\n")

SECTION_6 = [
    ("h", 1, "6. Possible Software Architecture"),
    ("h", 2, "6.1 Architectural Style"),
    ("p", "A monolithic full-stack Next.js application, not a microservices split — the right "
          "tradeoff for a solo developer (Journal Entries 30, 31). One deployable codebase handles "
          "server-rendered pages, API/server-action business logic, and scheduled jobs. All "
          "persistent state lives in one PostgreSQL database (Supabase), avoiding cross-service data "
          "consistency problems that would otherwise need to be solved from day one."),
    ("h", 2, "6.2 Component Diagram"),
    ("code", ARCHITECTURE_DIAGRAM),
    ("h", 2, "6.3 Data Flow (Example: A Review Crosses the Seal Vote Threshold)"),
    ("numbered", [
        "A user submits an upvote on a review via a Next.js Server Action.",
        "The Server Action writes the Vote row via Prisma, then recomputes that review's net score.",
        "If the review is not yet seal-provisional and its net score has just reached the configured quality threshold (Journal Entries 8, 29), the Server Action creates a provisional SealAward for that review, tagged 'granted_via: vote_threshold' (Journal Entries 6, 11).",
        "Every day, Vercel Cron triggers an API route that scans provisional SealAwards, checks whether each one's net score is still non-negative, increments or resets its streak counter, and converts any that have reached 30 consecutive days to permanent (Journal Entry 12).",
        "The Title's precomputed seal_count aggregate is updated whenever a SealAward's status changes, so browse/sort pages never need to compute this live (Journal Entry 39).",
    ]),
    ("h", 2, "6.4 Scheduled Jobs"),
    ("bullets", [
        "Daily probation-streak check (Vercel Cron → Next.js API route → Prisma) — the only scheduled job required for v1's seal system (Journal Entries 12, 36).",
        "Future: a periodic taste-profile recomputation job would be added when the recommendation engine (Section 2.9) is actually built — not required for v1.",
    ]),
    ("h", 2, "6.5 Scalability Strategy"),
    ("bullets", [
        "Vertical scaling of the single Postgres instance (Supabase paid tiers) covers a long runway before any architectural change is needed.",
        "Precomputed aggregate columns and targeted composite indexes keep hot-path reads fast as the catalog and review volume grow (Journal Entry 39).",
        "Read replicas are the identified next step if recommendation/analytics-style queries start competing with user-facing traffic — an infrastructure upgrade, not a re-platform.",
        "Storage layer (cover images) can be migrated independently to Cloudflare R2 if bandwidth becomes a cost driver, without touching the database or auth layers (Journal Entry 35).",
    ]),
]

DATA_MODEL_CODE = """
User
 - id, username, email, password_hash, avatar_url, bio, created_at,
   reputation_score, role (user | admin | main_admin),
   admin_reviews_require_approval (bool; only meaningful when role = admin)

Title
 - id, name, alt_names[], type (manga | manhwa | manhua), status,
   author, illustrator, genres[], cover_url, synopsis, publication_year,
   external_links[], created_at
 - precomputed: avg_category_scores, review_count,
   certified_banger_count, hidden_gem_count      <- Entry 39

Category                <- data-driven, not hardcoded         (Entry 2)
 - id, name, applies_to_type, scale_min, scale_max

Review
 - id, user_id, title_id, overall_score (= avg of category scores, Entry 4),
   body_text, spoiler_flag, created_at, updated_at,
   is_admin_authored (bool),
   approval_status (published | pending_approval | rejected)   <- Entry 28
   is_first_review_of_title (bool, derived)                    <- Entry 26

ReviewCategoryScore     <- one row per category per review
 - id, review_id, category_id, score

SealType                <- data-driven                        (Entry 14, 15)
 - id, name, description, icon
   (v1 rows: "Certified Banger", "Hidden Gem")

SealAward                                                      (Entry 5, 6, 7)
 - id, review_id, seal_type_id, justification_text,
   status (provisional | permanent),
   granted_via (admin | vote_threshold),
   granted_by_admin_id (nullable),
   granted_at, positive_streak_days, last_streak_reset_at,
   permanent_at (nullable), admin_override_status (nullable)   <- Entry 11-13

LibraryEntry             <- new v1 entity                     (Entry 18, 19)
 - id, user_id, title_id,
   status (finished | currently_reading | plan_to_read | dropped),
   added_at, updated_at

Comment
 - id, user_id, review_id, body_text, created_at

Vote
 - id, user_id, target_type (review | comment), target_id,
   value (up | down)
   UNIQUE (user_id, target_type, target_id)                    <- Section 4.5

Report
 - id, reporter_user_id, target_type, target_id, reason,
   status, created_at
""".strip("\n")

SECTION_7 = [
    ("h", 1, "7. Data Model Overview"),
    ("p", "Extends the brief's original data model sketch (Section 6 of README.md) with the "
          "decisions made in the Development Journal: admin role fields, the review-approval gate, "
          "the seal verification/probation state machine, and the new LibraryEntry entity. This is "
          "a design reference, not a final migration file — refine field types/constraints during "
          "implementation."),
    ("code", DATA_MODEL_CODE),
]

WORK_PACKAGES = [
    ("WP0.1", "Project scaffolding", "Next.js + TypeScript project init, Prisma setup, Supabase project (DB/Auth/Storage), Vercel deploy pipeline, environment/secrets configuration.", "—", "S"),
    ("WP0.2", "Core data model & migrations", "Implement all entities from Section 7; seed the Category table (5 rows) and SealType table (2 rows).", "WP0.1", "M"),
    ("WP1.1", "Authentication & user profiles", "Supabase Auth integration, registration/login, public profile page, role field (user/admin/main_admin).", "WP0.2", "M"),
    ("WP1.2", "Title catalog & admin seeding tools", "Admin panel to create/edit titles, search-before-create de-dup UX, cover upload, admin merge tooling.", "WP0.2", "M"),
    ("WP2.1", "Review creation", "5-category rating form, free-text body, computed overall score, spoiler flag, one-review-per-user-per-title with edit-replace.", "WP1.1, WP1.2", "M"),
    ("WP2.2", "Baseline review-submission gate", "Enforce email verification plus a minimum 3-day account age before first review submission; gate parameters admin-configurable for future tuning.", "WP2.1", "S"),
    ("WP2.3", "Admin review-approval workflow", "'(admin)' label everywhere a username appears; Main Admin approval queue for non-Main-Admin admin-authored reviews.", "WP2.1", "M"),
    ("WP2.4", "Title pages", "Aggregate score display, review list, precomputed-aggregate wiring (Journal Entry 39).", "WP2.1", "M"),
    ("WP3.1", "Voting on reviews", "Up/down voting with a unique constraint per user/review; net-score computation.", "WP2.1", "S"),
    ("WP3.2", "Comments on reviews", "Flat (non-threaded) comments.", "WP2.1", "S"),
    ("WP3.3", "Reporting/flagging", "Report flow for reviews, comments, and titles.", "WP2.1, WP3.2", "S"),
    ("WP4.1", "Seal data model & admin verification queue", "SealAward entity, justification field, Phase 1 admin-verifies-and-grants flow.", "WP2.1, WP3.1", "L"),
    ("WP4.2", "Vote-threshold auto-certification + probation cron", "Phase 2 automatic candidacy on threshold cross, Vercel Cron daily streak job, provisional→permanent conversion, manual admin override.", "WP4.1", "L"),
    ("WP4.3", "Hidden Gem popularity-gate logic", "Popularity threshold check, dual-seal (Certified Banger + Hidden Gem) coexistence and display.", "WP4.2", "M"),
    ("WP4.4", "Seal showcase page", "Dedicated Certified Banger / Hidden Gem discovery feed.", "WP4.1", "S"),
    ("WP5.1", "Search, filter & sort", "PostgreSQL full-text search, genre/status/score/seal filters, multi-field sort on the browse page.", "WP2.4, WP4.3", "M"),
    ("WP5.2", "User library feature", "Add/edit library entries (4 statuses), independent of reviews, displayed on profile.", "WP1.1, WP1.2", "M"),
    ("WP6.1", "Admin moderation panel", "Reports queue, ban/suspend users, title metadata editing, Category/SealType management UI.", "WP3.3, WP1.2", "M"),
    ("WP6.2", "Rate limiting & spam mitigation", "Rate limits on review/comment/vote/report creation.", "WP2.1, WP3.1, WP3.2, WP3.3", "S"),
    ("WP7.1", "Accessibility pass", "Alt text on covers, keyboard navigation, contrast check across all pages.", "All prior UI work packages", "S"),
    ("WP7.2", "Responsive/mobile QA", "Manual pass across core flows on mobile viewport sizes.", "All prior UI work packages", "S"),
    ("WP7.3", "Launch content seeding", "Main Admin authors initial reviews across the seeded title catalog.", "WP2.3, WP4.1", "M"),
]

SECTION_8 = [
    ("h", 1, "8. Work Packages & Build Sequence"),
    ("p", "Sequenced by dependency, not calendar date — as a solo project, actual duration depends "
          "on available hours per week, which wasn't specified. Relative size (S/M/L) is a rough "
          "effort indicator, not a time estimate. Work within a phase can generally proceed in any "
          "order; phases themselves are meant to be worked roughly top-to-bottom, since later phases "
          "depend on entities/flows built in earlier ones."),
    ("h", 2, "Phase 0 — Foundation & Infrastructure"),
    ("table", ["ID", "Work Package", "Description", "Depends On", "Size"],
     [(w[0], w[1], w[2], w[3], w[4]) for w in WORK_PACKAGES if w[0].startswith("WP0")]),
    ("h", 2, "Phase 1 — Identity & Catalog"),
    ("table", ["ID", "Work Package", "Description", "Depends On", "Size"],
     [(w[0], w[1], w[2], w[3], w[4]) for w in WORK_PACKAGES if w[0].startswith("WP1")]),
    ("h", 2, "Phase 2 — Core Review Loop"),
    ("table", ["ID", "Work Package", "Description", "Depends On", "Size"],
     [(w[0], w[1], w[2], w[3], w[4]) for w in WORK_PACKAGES if w[0].startswith("WP2")]),
    ("h", 2, "Phase 3 — Community Layer"),
    ("table", ["ID", "Work Package", "Description", "Depends On", "Size"],
     [(w[0], w[1], w[2], w[3], w[4]) for w in WORK_PACKAGES if w[0].startswith("WP3")]),
    ("h", 2, "Phase 4 — Seal System"),
    ("table", ["ID", "Work Package", "Description", "Depends On", "Size"],
     [(w[0], w[1], w[2], w[3], w[4]) for w in WORK_PACKAGES if w[0].startswith("WP4")]),
    ("h", 2, "Phase 5 — Discovery & Library"),
    ("table", ["ID", "Work Package", "Description", "Depends On", "Size"],
     [(w[0], w[1], w[2], w[3], w[4]) for w in WORK_PACKAGES if w[0].startswith("WP5")]),
    ("h", 2, "Phase 6 — Moderation & Abuse Mitigation"),
    ("table", ["ID", "Work Package", "Description", "Depends On", "Size"],
     [(w[0], w[1], w[2], w[3], w[4]) for w in WORK_PACKAGES if w[0].startswith("WP6")]),
    ("h", 2, "Phase 7 — Launch Prep"),
    ("table", ["ID", "Work Package", "Description", "Depends On", "Size"],
     [(w[0], w[1], w[2], w[3], w[4]) for w in WORK_PACKAGES if w[0].startswith("WP7")]),
    ("h", 2, "Post-v1 / Future Scope (Not Sequenced)"),
    ("bullets", [
        "F1 — Recommendation engine (Journal Entries 20, 21, 24, 25).",
        "F2 — First-review token reward (Journal Entries 22, 26).",
        "F3 — Token/crypto integration (Brief Section 8).",
        "F4 — Multi-media expansion, public API, reviewer reputation/leveling, native mobile app (Brief Section 8).",
    ]),
]

USER_STORIES = [
    ("Authentication & Profiles", [
        "As a visitor, I want to browse titles and read reviews without creating an account, so I can evaluate the platform's value before signing up.",
        "As a visitor, I want to register with an email and password, so I can start writing reviews and voting.",
        "As a registered user, I want to log in and out, so I can access my account securely.",
        "As a registered user, I want a public profile page showing my reviews, seals awarded, join date, and library, so others can gauge my track record.",
        "As any user, I want an admin's username to always show an '(admin)' label wherever it appears, so I can tell official content apart from community content.",
    ]),
    ("Title Catalog", [
        "As an admin, I want to create and edit title entries (name, alt names, type, authors, status, genres, cover, synopsis, source links), so the catalog stays accurate and complete.",
        "As an admin, I want a search-before-create flow when adding a title, so duplicate entries aren't created.",
        "As an admin, I want to merge duplicate title entries, so the catalog stays canonical.",
        "As an admin, I want rating categories and seal types to be configurable through an admin interface, so I can add a 6th category or a new seal type without a code deploy.",
        "As a user, I want to view a title's page showing its aggregate category scores, seal counts, and review list, so I can evaluate it at a glance.",
    ]),
    ("Reviews", [
        "As a registered user, I want to write one review per title with a rating for each of the five categories, so my opinion is structured and comparable to others'.",
        "As a registered user, I want my review's overall score computed automatically as the average of my category scores, so I don't have to separately judge a single number.",
        "As a registered user, I want to mark my review as containing spoilers, so readers can choose whether to see the full text.",
        "As a registered user, I want to edit or delete my own review, so I can correct or update my opinion later.",
        "As a new user, I want to understand the minimum account age required before I can submit a review, so I know why I might be blocked.",
        "As the Main Admin, I want my own reviews to publish immediately without needing approval, so I can seed content efficiently at launch.",
        "As an admin who is not the Main Admin, I want my reviews to enter a pending-approval queue, so the Main Admin can verify quality before they go live.",
        "As the Main Admin, I want to approve or reject pending admin-authored reviews, so I can maintain quality control over official content.",
    ]),
    ("Seals — Certified Banger & Hidden Gem", [
        "As a registered user, I want to see whether a review has earned the Certified Banger or Hidden Gem seal, so I can trust the platform's curation signal.",
        "As a reviewer, I want to write a justification separate from my review body when nominating a title for a seal, so my rating rationale and my 'why this deserves recognition' argument stay distinct.",
        "As the Main Admin, I want to manually verify a submitted review and grant it a seal, so I can curate quality seals while the platform is new.",
        "As a registered user, I want a review's votes to automatically qualify it for a seal once it crosses the configured vote threshold, so seal-granting scales without an admin bottleneck once the community is large enough.",
        "As a user, I want a newly-earned automatic seal to stay 'provisional' until it holds a positive vote streak for 30 consecutive days, so short-term vote brigading can't cheapen the seal.",
        "As a user, I want a provisional seal to convert to permanent automatically once it achieves the 30-day streak, so genuinely good reviews get recognized without manual intervention.",
        "As the Main Admin, I want to manually demote or remove a permanent seal, so I can correct fraud or abuse discovered after the fact.",
        "As a registered user, I want a title with a high-quality but low-popularity review to earn the Hidden Gem seal, so I can discover well-reviewed titles that aren't yet widely known.",
        "As a registered user, I want a title to keep its Hidden Gem seal even after it also earns Certified Banger, so the 'this was an underrated find' signal isn't lost once a title becomes popular.",
    ]),
    ("Voting & Comments", [
        "As a registered user, I want to upvote or downvote a review, so the community can surface the most helpful/agreed-with reviews.",
        "As a registered user, I want to comment on a review, so I can discuss or disagree with it.",
        "As a user, I want to see engagement (votes, comment count) on a review, so I can judge how the community responded to it.",
    ]),
    ("Discovery & Browsing", [
        "As a user, I want to search titles by name, genre, or author, so I can quickly find what I'm looking for.",
        "As a user, I want to filter titles by genre, status, category-score thresholds, and seal presence, so I can narrow down to what matters to me.",
        "As a user, I want to sort titles by highest overall score, most seals, or most recent reviews, so I can browse in the order that's useful to me.",
        "As a user, I want a dedicated Certified Banger / Hidden Gem showcase page, so the platform's curated highlights are easy to find.",
    ]),
    ("User Library", [
        "As a registered user, I want to add a title to my library with a status (Finished, Currently Reading, Plan to Read, or Dropped), so I can track my reading progress.",
        "As a registered user, I want to add a title to my library without needing to review it, so I can track titles I haven't formed a full opinion on yet.",
        "As a registered user, I want to change a library entry's status as I progress (e.g. Currently Reading to Finished), so my library stays accurate.",
        "As a registered user, I want to view my library on my profile, so I (and others) can see what I've read or am reading.",
    ]),
    ("Moderation & Reporting", [
        "As a registered user, I want to report a review, comment, or title as inappropriate/spam, so problematic content gets moderator attention.",
        "As an admin, I want a queue of reported content, so I can review and act on flags efficiently.",
        "As an admin, I want to remove content or ban/suspend a user, so I can enforce platform standards.",
        "As a platform, I want rate limits on review, comment, and report creation, so spam and abuse are mitigated automatically.",
    ]),
    ("Future Scope (Backlog, Not v1)", [
        "As a registered user, I want personalized title recommendations based on my finished library's category ratings and genres, so I can discover new titles suited to my taste. (Algorithm undesigned — Journal Entry 21.)",
        "As a registered user, I want to receive a token reward if my review is the first, well-received review of a title, so I'm incentivized to help seed review coverage across the catalog. (Journal Entries 22, 26.)",
    ]),
]

SECTION_9 = [("h", 1, "9. User Stories")]
SECTION_9.append(("p", "Organized by epic, matching the functional-requirement areas in Section 2 "
                       "and the work packages in Section 8, so stories can be picked up work-package "
                       "by work-package."))
for epic_name, stories in USER_STORIES:
    SECTION_9.append(("h", 2, epic_name))
    SECTION_9.append(("bullets", stories))

SECTION_10 = [
    ("h", 1, "10. Open Items / Known Unknowns"),
    ("p", "As of this revision, all v1-scoped open items have been resolved (Journal Entries 40–42). "
          "The remaining open items below are future-scope details that were explicitly confirmed to "
          "stay deferred by design — not oversights — until the feature they belong to enters active "
          "development. None of them block any Section 8 work package."),
    ("table", ["Open Item", "Affects", "Journal Reference"], [
        ("Recommendation engine algorithm design (content-based filtering, hybrid, or other).", "Future Scope F1", "Entry 21, 23"),
        ("Whether 'Dropped' library status becomes a negative signal for recommendations in a future iteration.", "Future Scope F1", "Entry 21, 25"),
        ("Exact 'well received' vote threshold for the first-review token reward.", "Future Scope F2", "Entry 22, 23"),
    ]),
    ("h", 2, "Resolved Since Previous Revision"),
    ("bullets", [
        "Baseline review-submission gate — exact criteria and evolution process (Journal Entry 40).",
        "Admin trust criteria for lifting the review-approval requirement — manual case-by-case (Journal Entry 41).",
        "Confirmed: unapproved admin reviews are fully hidden until approved (Journal Entry 42).",
    ]),
]


def main():
    document = Document()
    section = document.sections[0]
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    document.styles["Normal"].font.size = Pt(11)

    build_styles(document)
    enable_auto_update_fields(document)

    add_title_page(document)
    add_toc(document)

    render_blocks(document, SECTION_1)
    document.add_page_break()
    render_blocks(document, SECTION_2)
    document.add_page_break()
    render_blocks(document, SECTION_3)
    document.add_page_break()
    render_blocks(document, SECTION_4)
    document.add_page_break()
    render_blocks(document, SECTION_5)
    document.add_page_break()
    render_blocks(document, SECTION_6)
    document.add_page_break()
    render_blocks(document, SECTION_7)
    document.add_page_break()
    render_blocks(document, SECTION_8)
    document.add_page_break()
    render_blocks(document, SECTION_9)
    document.add_page_break()
    render_blocks(document, SECTION_10)

    document.save(OUTPUT_PATH)
    print(f"Wrote {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
