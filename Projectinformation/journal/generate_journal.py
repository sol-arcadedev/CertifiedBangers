"""
Regenerates Development_Journal.docx from entries.json.

Usage: python generate_journal.py

Edit entries.json to add/update journal entries, then rerun this script
to rebuild the docx. Re-running always overwrites the output file.
"""
import json
import os
from datetime import datetime

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = os.path.dirname(os.path.abspath(__file__))
ENTRIES_PATH = os.path.join(HERE, "entries.json")
OUTPUT_PATH = os.path.join(os.path.dirname(HERE), "Development_Journal.docx")

STATUS_COLORS = {
    "Completed": RGBColor(0x2E, 0x7D, 0x32),
    "Resolved": RGBColor(0x2E, 0x7D, 0x32),
    "Open": RGBColor(0xC6, 0x28, 0x28),
    "Open Question Log": RGBColor(0xC6, 0x28, 0x28),
}


def status_color(status: str) -> RGBColor:
    for key, color in STATUS_COLORS.items():
        if status.startswith(key):
            return color
    return RGBColor(0xE6, 0x8A, 0x00)  # amber for partial/mixed statuses


def add_field(paragraph, instr_text, display_text):
    run = paragraph.add_run()
    r = run._r

    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")

    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instr_text

    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")

    text_el = OxmlElement("w:t")
    text_el.text = display_text

    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")

    r.append(fld_begin)
    r.append(instr)
    r.append(fld_sep)
    r.append(text_el)
    r.append(fld_end)


def enable_auto_update_fields(document):
    settings = document.settings.element
    update_fields = OxmlElement("w:updateFields")
    update_fields.set(qn("w:val"), "true")
    settings.append(update_fields)


def set_cell_shading(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def build_styles(document):
    styles = document.styles

    title_style = styles.add_style("JournalTitle", WD_STYLE_TYPE.PARAGRAPH)
    title_style.base_style = styles["Title"]
    title_style.font.size = Pt(28)
    title_style.font.color.rgb = RGBColor(0x1A, 0x1A, 0x2E)

    meta_style = styles.add_style("Meta", WD_STYLE_TYPE.PARAGRAPH)
    meta_style.font.size = Pt(10)
    meta_style.font.color.rgb = RGBColor(0x60, 0x60, 0x60)
    meta_style.font.italic = True

    label_style = styles.add_style("FieldLabel", WD_STYLE_TYPE.CHARACTER)
    label_style.font.bold = True
    label_style.font.size = Pt(11)
    label_style.font.color.rgb = RGBColor(0x1A, 0x1A, 0x2E)


def add_title_page(document):
    p = document.add_paragraph(style="JournalTitle")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("CertifiedBanger")

    sub = document.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = sub.add_run("Project Development Journal")
    run.font.size = Pt(18)
    run.font.color.rgb = RGBColor(0x40, 0x40, 0x40)

    meta = document.add_paragraph(style="Meta")
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.add_run(f"Generated {datetime.now().strftime('%Y-%m-%d')}")

    document.add_paragraph()
    purpose_heading = document.add_paragraph()
    purpose_heading.add_run("Purpose of this document").bold = True

    purpose_body = document.add_paragraph(
        "This journal records every substantive decision, question, and open item that arises "
        "during the design and development of CertifiedBanger. Each entry documents: what the "
        "step or decision was about, what options were considered, which option was chosen, and "
        "why — so that any decision can be reviewed and traced back to its reasoning later. "
        "Open items are logged even when unresolved, and are updated in place once a decision is made."
    )
    purpose_body.paragraph_format.space_after = Pt(12)

    legend_heading = document.add_paragraph()
    legend_heading.add_run("Status legend").bold = True

    legend_table = document.add_table(rows=1, cols=2)
    legend_table.style = "Light Grid Accent 1"
    hdr = legend_table.rows[0].cells
    hdr[0].text = "Status"
    hdr[1].text = "Meaning"

    legend_rows = [
        ("Completed / Resolved", "The step is finished, or the decision has been made and is considered settled."),
        ("Partially Resolved", "A core decision has been made, but one or more sub-questions remain open."),
        ("Open", "Not yet decided — awaiting input or further discussion."),
        ("Superseded", "An earlier decision that has since been replaced or substantially refined by a later entry."),
    ]
    for status, meaning in legend_rows:
        row = legend_table.add_row().cells
        row[0].text = status
        row[1].text = meaning

    document.add_page_break()


def add_toc(document):
    heading = document.add_heading("Table of Contents", level=1)
    heading.paragraph_format.space_after = Pt(12)

    toc_paragraph = document.add_paragraph()
    add_field(
        toc_paragraph,
        'TOC \\o "1-2" \\h \\z \\u',
        "Table of Contents will appear here — in Word, right-click and choose "
        "'Update Field', or press Ctrl+A then F9, to generate it.",
    )
    document.add_page_break()


def add_key_value(document, label, value, color=None):
    p = document.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    run_label = p.add_run(f"{label}: ")
    run_label.style = document.styles["FieldLabel"]
    run_value = p.add_run(str(value))
    if color:
        run_value.font.color.rgb = color
        run_value.font.bold = True


def add_options_table(document, options):
    if not options:
        return
    table = document.add_table(rows=1, cols=2)
    table.style = "Light Grid Accent 1"
    table.autofit = True
    hdr = table.rows[0].cells
    hdr[0].text = "Option Considered"
    hdr[1].text = "Description"
    for cell in hdr:
        for para in cell.paragraphs:
            for run in para.runs:
                run.bold = True
    for opt in options:
        row = table.add_row().cells
        row[0].text = opt.get("name", "")
        row[1].text = opt.get("description", "")
    document.add_paragraph().paragraph_format.space_after = Pt(6)


def add_open_items_table(document, items):
    table = document.add_table(rows=1, cols=2)
    table.style = "Light Grid Accent 2"
    hdr = table.rows[0].cells
    hdr[0].text = "Topic"
    hdr[1].text = "Open Question"
    for cell in hdr:
        for para in cell.paragraphs:
            for run in para.runs:
                run.bold = True
    for item in items:
        row = table.add_row().cells
        row[0].text = item.get("topic", "")
        row[1].text = item.get("question", "")


def add_entry(document, entry):
    heading = document.add_heading(f"Entry {entry['id']}: {entry['title']}", level=1)
    heading.paragraph_format.space_before = Pt(18)

    add_key_value(document, "Date", entry.get("date", ""))
    add_key_value(document, "Type", entry.get("type", ""))
    add_key_value(document, "Status", entry.get("status", ""), color=status_color(entry.get("status", "")))
    if entry.get("relatesTo"):
        add_key_value(document, "Relates To", entry["relatesTo"])

    document.add_heading("What this step was about", level=2)
    document.add_paragraph(entry.get("summary", ""))
    if entry.get("context"):
        document.add_paragraph(entry["context"])

    if entry.get("options"):
        document.add_heading("Options considered", level=2)
        add_options_table(document, entry["options"])

    if entry.get("decision"):
        document.add_heading("Decision made", level=2)
        document.add_paragraph(entry["decision"])

    if entry.get("rationale"):
        document.add_heading("Why this decision was made", level=2)
        document.add_paragraph(entry["rationale"])

    if entry.get("openSubItems"):
        document.add_heading("Open sub-items", level=2)
        for item in entry["openSubItems"]:
            document.add_paragraph(item, style="List Bullet")

    if entry.get("items"):
        document.add_heading("Open questions", level=2)
        add_open_items_table(document, entry["items"])

    document.add_paragraph()


def main():
    with open(ENTRIES_PATH, "r", encoding="utf-8") as f:
        entries = json.load(f)

    document = Document()

    section = document.sections[0]
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    style = document.styles["Normal"]
    style.font.size = Pt(11)

    build_styles(document)
    enable_auto_update_fields(document)

    add_title_page(document)
    add_toc(document)

    for entry in entries:
        add_entry(document, entry)

    document.save(OUTPUT_PATH)
    print(f"Wrote {OUTPUT_PATH} ({len(entries)} entries)")


if __name__ == "__main__":
    main()
